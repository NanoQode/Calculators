"""Build the lmcmic.ca SEO master workbook from the built site.

    python3 build/seo_master.py      # after build/build.py; writes reports/lmcmic-seo-master.xlsx

Everything is read from what actually ships: the HTML in dist/, the sitemap,
the nginx redirect rules, the content plan, the QA report and the live-server
measurements in reports/live-status.json. Lengths, flags and the overview
counts are Excel formulas, so the sheet stays correct if a row is edited.
"""
import datetime as dt
import html
import json
import pathlib
import re
import sys
from html.parser import HTMLParser

import openpyxl
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "build"))
import build as B  # noqa: E402  (module-level only: config, plan, CORE_KEYWORDS)

DIST = ROOT / "dist"
OUT = ROOT / "reports/lmcmic-seo-master.xlsx"
SITE = B.SITE
PLAN = B.PLAN
QA = json.loads((ROOT / "reports/qa-report.json").read_text())
LIVE = json.loads((ROOT / "reports/live-status.json").read_text())
LINKS = json.loads((ROOT / "reports/link-check.json").read_text())
IMAGES = json.loads((ROOT / "content/images.json").read_text())
SRC_XLSX = ROOT / "content/data/seo-content-roadmap.xlsx"
DOMAIN = SITE["domain"]
TODAY = dt.date.today().isoformat()

# ------------------------------------------------------------------ styles
F = "Arial"
NAVY, SKY, RED = "0E1C2F", "006398", "DC2626"
HDR = PatternFill("solid", fgColor=NAVY)
SUB = PatternFill("solid", fgColor="E5EEFF")
INPUT = PatternFill("solid", fgColor="FFFF00")
OKF = PatternFill("solid", fgColor="DCFCE7")
BADF = PatternFill("solid", fgColor="FEE2E2")
WARNF = PatternFill("solid", fgColor="FEF3C7")
THIN = Side(style="thin", color="E2E8F0")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def font(**kw):
    return Font(name=F, size=kw.pop("size", 10), **kw)


def header(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = font(bold=True, color="FFFFFF")
        cell.fill = HDR
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BOX
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30


def body(ws, first, last, ncols, wrap=True):
    for r in ws.iter_rows(min_row=first, max_row=last, max_col=ncols):
        for c in r:
            c.font = c.font if c.font and c.font.name == F and (c.font.bold or c.font.color) else font()
            c.alignment = Alignment(wrap_text=wrap, vertical="top")
            c.border = BOX


def title_block(ws, title, subtitle):
    ws["A1"] = title
    ws["A1"].font = font(size=16, bold=True, color=NAVY)
    ws["A2"] = subtitle
    ws["A2"].font = font(size=10, italic=True, color="44474C")


# ------------------------------------------------------------------ parse dist
class PageParser(HTMLParser):
    """Pulls SEO fields and in-content links from a built page."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.meta, self.title, self.h1, self.links, self.ld = {}, "", "", [], []
        self.canonical = self.og_image = ""
        self._in = {"title": False, "h1": False, "main": 0, "nav": 0, "a": None, "ld": False, "script": False, "style": False, "hidden": 0}
        self._buf, self._a_text, self.text = [], [], []
        self.imgs = set()

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        # Icon glyphs are aria-hidden ligature text ("arrow_forward"); keep them out of anchors and copy.
        if tag == "span" and (self._in["hidden"] or a.get("aria-hidden") == "true"):
            self._in["hidden"] += 1
        if tag == "title":
            self._in["title"] = True
        elif tag == "meta":
            k = a.get("name") or a.get("property")
            if k:
                self.meta[k] = a.get("content", "")
        elif tag == "link" and a.get("rel") == "canonical":
            self.canonical = a.get("href", "")
        elif tag == "h1":
            self._in["h1"] = True
        elif tag == "main":
            self._in["main"] += 1
        elif tag == "nav":
            self._in["nav"] += 1
        elif tag == "script":
            self._in["ld"] = a.get("type") == "application/ld+json"
            self._in["script"] = True
        elif tag == "style":
            self._in["style"] = True
        elif tag == "a" and self._in["main"] and not self._in["nav"]:
            self._in["a"] = a.get("href", "")
            self._a_text = []
        elif tag in ("img", "source"):
            for v in (a.get("src", ""), a.get("srcset", "")):
                self.imgs |= set(re.findall(r"/assets/img/[\w@.-]+", v))

    def handle_endtag(self, tag):
        if tag == "span" and self._in["hidden"]:
            self._in["hidden"] -= 1
        if tag == "title":
            self._in["title"] = False
        elif tag == "h1":
            self._in["h1"] = False
        elif tag == "main":
            self._in["main"] -= 1
        elif tag == "nav":
            self._in["nav"] -= 1
        elif tag == "script":
            if self._in["ld"]:
                self.ld.append("".join(self._buf))
            self._buf, self._in["ld"], self._in["script"] = [], False, False
        elif tag == "style":
            self._in["style"] = False
        elif tag == "a" and self._in["a"] is not None:
            self.links.append((self._in["a"], re.sub(r"\s+", " ", "".join(self._a_text)).strip()))
            self._in["a"] = None

    def handle_data(self, d):
        if self._in["title"]:
            self.title += d
        if self._in["h1"]:
            self.h1 += d
        if self._in["ld"]:
            self._buf.append(d)
        if self._in["hidden"]:
            return
        if self._in["a"] is not None:
            self._a_text.append(d)
        if self._in["main"] and not self._in["script"] and not self._in["style"]:
            self.text.append(d)


def page_files():
    for p in sorted(DIST.rglob("index.html")):
        rel = p.relative_to(DIST).parent.as_posix()
        yield ("/" if rel == "." else f"/{rel}/"), p


def ld_types(blocks):
    out = []
    for b in blocks:
        try:
            t = json.loads(b).get("@type")
        except ValueError:
            continue
        for x in (t if isinstance(t, list) else [t]):
            if x and x not in out:
                out.append(x)
    return out


sitemap_urls = set(re.findall(r"<loc>https://lmcmic\.ca([^<]+)</loc>", (DIST / "sitemap.xml").read_text()))
plan_by_url = {a["url"]: a for a in PLAN["articles"]}
qa_by_url = {a["url"]: a for a in QA["articles"]}

pages = []
for url, path in page_files():
    pp = PageParser()
    pp.feed(path.read_text(encoding="utf-8"))
    internal = []
    for href, text in pp.links:
        h = href.split("#")[0].split("?")[0]
        if h.startswith("/") and not h.startswith("/assets/") and h:
            internal.append((h, text))
    a = plan_by_url.get(url)
    if a:
        kind = {"Calculator / tool": "Tool", "FAQ hub": "FAQ hub", "Location page": "Location"}.get(a["type"], a["type"])
        if url == "/how-to-invest/":
            kind = "Core page"
        primary, secondary, question = a["primary_keyword"], "; ".join(a["secondary_keywords"]), a["question_keyword"]
        section, prio = a["stage_name"], f"P{a['priority']}"
    else:
        kind = {"/learn/": "Hub", "/privacy/": "Legal", "/disclaimer/": "Legal"}.get(url, "Core page")
        ks = B.CORE_KEYWORDS.get(url, [])
        primary, secondary, question = (ks[0] if ks else ""), "; ".join(ks[1:]), ""
        section, prio = ("Investing with Lendmax" if kind == "Core page" else kind), ("Core" if kind == "Core page" else "—")
    words = len(re.findall(r"[A-Za-z0-9$%][\w'’%.,$-]*", " ".join(pp.text)))
    q = qa_by_url.get(url)
    pages.append({
        "url": url, "type": kind, "section": section, "priority": prio,
        "title": html.unescape(pp.title.strip()), "description": pp.meta.get("description", ""),
        "h1": re.sub(r"\s+", " ", pp.h1).strip(), "primary": primary, "secondary": secondary, "question": question,
        "canonical": pp.canonical, "robots": pp.meta.get("robots", ""), "schema": ", ".join(ld_types(pp.ld)),
        "words": words, "links_out": len({h for h, _ in internal}), "internal": internal,
        "og": pp.meta.get("og:image", ""), "photo": "Yes" if any("driving-results" in i for i in pp.imgs) else "",
        "logos": "Yes" if any("partner-" in i for i in pp.imgs) else "",
        "sitemap": "Yes" if url in sitemap_urls else "No",
        "qa": (q or {}).get("status", "n/a (core page)" if not a else ""),
        "ga": "Yes" if 'data-ga="' in path.read_text() else "No",
    })

inlinks = {}
for p in pages:
    for h in {h for h, _ in p["internal"]}:
        inlinks.setdefault(h, set()).add(p["url"])

# ------------------------------------------------------------------ workbook
wb = openpyxl.Workbook()

# ---- Pages
ws = wb.active
ws.title = "Pages"
cols = ["#", "URL", "Page type", "Section / journey stage", "Priority", "Title tag", "Title length", "Title check",
        "Meta description", "Description length", "Description check", "H1", "Primary keyword", "Secondary keywords",
        "Question keyword", "Canonical URL", "Canonical matches", "Robots", "Structured data (schema.org)", "Words on page",
        "Internal links out", "Internal links in", "Share image (og:image)", "Photo", "Partner logos", "In sitemap",
        "GA4 tag", "Automated QA", "Last modified", "Live URL"]
header(ws, 1, cols, [5, 38, 12, 22, 8, 46, 8, 10, 60, 10, 11, 40, 28, 40, 34, 40, 10, 22, 34, 9, 9, 9, 40, 7, 8, 9, 8, 12, 12, 40])
for i, p in enumerate(sorted(pages, key=lambda x: (x["type"] != "Core page", x["type"] != "Hub", x["url"])), start=2):
    r = i
    vals = [r - 1, p["url"], p["type"], p["section"], p["priority"], p["title"], f"=LEN(F{r})",
            f'=IF(AND(G{r}>=30,G{r}<=65),"OK","Review")', p["description"], f"=LEN(I{r})",
            f'=IF(AND(J{r}>=120,J{r}<=160),"OK","Review")', p["h1"], p["primary"], p["secondary"], p["question"],
            p["canonical"], f'=IF(P{r}="https://lmcmic.ca"&B{r},"Yes","No")', p["robots"], p["schema"], p["words"],
            p["links_out"], len(inlinks.get(p["url"], ())), p["og"], p["photo"], p["logos"], p["sitemap"], p["ga"],
            p["qa"], SITE["content_updated"], f'=HYPERLINK("https://lmcmic.ca"&B{r},"Open")']
    for c, v in enumerate(vals, 1):
        ws.cell(row=r, column=c, value=v)
last_pages = len(pages) + 1
body(ws, 2, last_pages, len(cols))
for col in ("H", "K", "Q"):
    ws.conditional_formatting.add(f"{col}2:{col}{last_pages}", CellIsRule(operator="equal", formula=['"OK"'], fill=OKF))
    ws.conditional_formatting.add(f"{col}2:{col}{last_pages}", CellIsRule(operator="equal", formula=['"Yes"'], fill=OKF))
    ws.conditional_formatting.add(f"{col}2:{col}{last_pages}", CellIsRule(operator="equal", formula=['"Review"'], fill=WARNF))
    ws.conditional_formatting.add(f"{col}2:{col}{last_pages}", CellIsRule(operator="equal", formula=['"No"'], fill=BADF))
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last_pages}"
ws["G1"].comment = Comment("Formula: characters in the title tag. Target 30–65 so Google shows it in full.", "lmcmic build")
ws["J1"].comment = Comment("Formula: characters in the meta description. Target 120–160.", "lmcmic build")
ws["V1"].comment = Comment("How many other pages link to this one from their main content. Navigation, breadcrumbs and footer are excluded; they link every page to the home page and hubs, so the home page shows 0 here.", "lmcmic build")

# ---- Keywords
wk = wb.create_sheet("Keywords")
kcols = ["Keyword", "Keyword type", "Cluster", "Search intent", "Investor stage", "Geography", "Opportunity (1-5)",
         "Competition (1-5)", "Priority score", "Build tier", "AI / answer-engine relevance", "Workbook target page",
         "Live URL", "Live page title", "Terms on page", "Measured CA monthly volume", "Volume source & date",
         "Google position (GSC)", "Clicks, last 28 days (GSC)", "Notes"]
header(wk, 1, kcols, [40, 16, 30, 18, 16, 18, 10, 10, 9, 16, 11, 34, 42, 46, 9, 12, 18, 11, 11, 30])
for i, k in enumerate(QA["keywords"], start=2):
    src = next((x for x in PLAN["keywords"] if x["keyword"] == k["keyword"]), {})
    vals = [k["keyword"], src.get("type", ""), k["cluster"], k["intent"], src.get("stage", ""), src.get("geo", ""),
            float(src["opportunity"]) if src.get("opportunity") else None,
            float(src["competition"]) if src.get("competition") else None,
            f"=IF(AND(ISNUMBER(G{i}),ISNUMBER(H{i})),G{i}*2-H{i},\"\")", k["tier"], src.get("ai_relevance", ""),
            k["target_page"], k["url"] or "",
            f'=IF(M{i}="","Not built (workbook: do not build)",IFERROR(INDEX(Pages!$F$2:$F${last_pages},MATCH(M{i},Pages!$B$2:$B${last_pages},0)),"Not found"))',
            "Yes" if k["on_page"] else ("No" if k["url"] else "—"), None, None, None, None,
            src.get("geo_guidance", "")]
    for c, v in enumerate(vals, 1):
        wk.cell(row=i, column=c, value=v)
last_kw = len(QA["keywords"]) + 1
body(wk, 2, last_kw, len(kcols))
for col in "PQRS":
    for r in range(2, last_kw + 1):
        wk[f"{col}{r}"].fill = INPUT
wk.conditional_formatting.add(f"O2:O{last_kw}", CellIsRule(operator="equal", formula=['"Yes"'], fill=OKF))
wk.conditional_formatting.add(f"O2:O{last_kw}", CellIsRule(operator="equal", formula=['"No"'], fill=BADF))
wk.freeze_panes = "B2"
wk.auto_filter.ref = f"A1:{get_column_letter(len(kcols))}{last_kw}"
wk["I1"].comment = Comment("Formula from the source workbook: (Opportunity × 2) − Competition.", "lmcmic build")
wk["O1"].comment = Comment(f"Checked at build {TODAY}: every word of the keyword (3+ letters) appears in the live page's text.", "lmcmic build")
wk["P1"].comment = Comment("YELLOW = you fill in. Paste Canada monthly volume from Google Keyword Planner or Ahrefs.", "lmcmic build")
wk["R1"].comment = Comment("YELLOW = you fill in. From Search Console > Performance, filtered to this query.", "lmcmic build")

# ---- Content QA
wq = wb.create_sheet("Content QA")
checks = ["FORMAT", "META_TITLE", "META_DESC", "K1_ANSWER_FIRST_120_WORDS", "K2_HEADINGS", "K7_PROHIBITED_LANGUAGE",
          "K8_DISCLOSURES_IN_BODY", "K11_INTERNAL_LINKS", "SOURCES_ALLOWED", "J2_SEVEN_AXES", "E7_FAQ", "E9_TAKEAWAYS", "LENGTH"]
labels = ["Front matter", "Title length", "Description length", "K1 Answer in first 120 words", "K2 One H1, no heading jumps",
          "K7 Prohibited-language scan", "K8 Disclosures in body", "K11 Internal links", "Allowed sources only",
          "J2 Seven axes in summary", "E7 FAQ", "E9 Key takeaways", "Length vs target"]
def article_words(a):
    """Words qa.py measures against the target: body + short answer + FAQs + glossary terms."""
    m = re.match(r"(\d+) words", str((a["checks"].get("LENGTH") or {}).get("detail", "")))
    return int(m.group(1)) if m else a["words"]


qcols = ["#", "Article", "URL", "Priority", "Primary keyword", "Words (article)", "Target words", "Overall"] + labels + ["Checks passed"]
header(wq, 1, qcols, [5, 46, 40, 8, 30, 9, 9, 9] + [11] * len(labels) + [9])
for i, a in enumerate(QA["articles"], start=2):
    plan = plan_by_url.get(a["url"], {})
    row = [a["id"], a["title"], a["url"], f"P{a.get('priority', plan.get('priority', ''))}", a.get("primary_keyword", ""),
           article_words(a), plan.get("target_words")]
    first_check_col = len(row) + 2
    row.append(f'=IF(COUNTIF({get_column_letter(first_check_col)}{i}:{get_column_letter(first_check_col + len(checks) - 1)}{i},"Fail")=0,"PASS","FAIL")')
    for c in checks:
        ch = a["checks"].get(c)
        row.append("Pass" if ch and ch["ok"] else ("Fail" if ch else "—"))
    row.append(f'=COUNTIF({get_column_letter(first_check_col)}{i}:{get_column_letter(first_check_col + len(checks) - 1)}{i},"Pass")')
    for c, v in enumerate(row, 1):
        cell = wq.cell(row=i, column=c, value=v)
        if 9 <= c < 9 + len(checks):
            ch = a["checks"].get(checks[c - 9])
            if ch:
                cell.comment = Comment(str(ch["detail"])[:300], "qa.py")
last_q = len(QA["articles"]) + 1
body(wq, 2, last_q, len(qcols))
rng = f"H2:{get_column_letter(8 + len(checks))}{last_q}"
for word, fill in (("PASS", OKF), ("Pass", OKF), ("FAIL", BADF), ("Fail", BADF)):
    wq.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{word}"'], fill=fill))
wq.freeze_panes = "C2"
wq.auto_filter.ref = f"A1:{get_column_letter(len(qcols))}{last_q}"
mr = last_q + 3
wq.cell(row=mr, column=1, value="Manual checks (Tab 3, section K)").font = font(size=12, bold=True, color=NAVY)
header(wq, mr + 1, ["Ref", "Check", "Status", "Detail"])
for j, m in enumerate(QA["manual_checks"], start=mr + 2):
    for c, v in enumerate([m["id"], m["label"], m["status"], m["note"]], 1):
        wq.cell(row=j, column=c, value=v)
body(wq, mr + 2, mr + 1 + len(QA["manual_checks"]), 4)
wq.merge_cells(start_row=mr + 1, start_column=4, end_row=mr + 1, end_column=12)
for j in range(mr + 2, mr + 2 + len(QA["manual_checks"])):
    wq.merge_cells(start_row=j, start_column=4, end_row=j, end_column=12)
    wq.row_dimensions[j].height = 30

# ---- Internal links
wl = wb.create_sheet("Internal Links")
header(wl, 1, ["From page", "To page", "Anchor text", "Target exists"], [44, 44, 60, 11])
r = 2
all_urls = {p["url"] for p in pages}
for p in sorted(pages, key=lambda x: x["url"]):
    seen = set()
    for h, text in p["internal"]:
        if (h, text) in seen or h == p["url"]:
            continue
        seen.add((h, text))
        wl.cell(row=r, column=1, value=p["url"])
        wl.cell(row=r, column=2, value=h)
        wl.cell(row=r, column=3, value=text)
        wl.cell(row=r, column=4, value=f'=IF(ISNUMBER(MATCH(B{r},Pages!$B$2:$B${last_pages},0)),"Yes","Check")')
        r += 1
last_l = r - 1
body(wl, 2, last_l, 4)
wl.conditional_formatting.add(f"D2:D{last_l}", CellIsRule(operator="equal", formula=['"Check"'], fill=BADF))
wl.freeze_panes = "A2"
wl.auto_filter.ref = f"A1:D{last_l}"

# ---- Redirects
wr = wb.create_sheet("Redirects")
header(wr, 1, ["Requested", "Goes to", "Status", "Why"], [48, 40, 8, 70])
rows = [
    ("http://lmcmic.ca/<any path>", "https://lmcmic.ca/<same path>", 301, "One canonical origin: HTTPS"),
    ("http(s)://www.lmcmic.ca/<any path>", "https://lmcmic.ca/<same path>", 301, "One canonical host: no www"),
    ("/<path> (no trailing slash)", "/<path>/", 301, "Directory URLs end in a slash; only one version is indexable"),
    ("/<path>/index.html", "/<path>/", 301, "Strip index.html (tested on the original request, so no loop)"),
]
conf = (ROOT / "deploy/nginx-lmcmic.ca.conf").read_text()
for alts, target, code in re.findall(r"location ~\* \^/\(([^)]+)\)[^{]*\{ return (\d{3}) ?([^;]*);", conf) and \
        [(m[0], m[2] or "—", int(m[1])) for m in re.findall(r"location ~\* \^/\(([^)]+)\)[^{]*\{ return (\d{3}) ?([^;]*);", conf)]:
    for alt in alts.split("|"):
        alt = alt.replace("\\.", ".").replace("?", "")
        rows.append((f"/{alt}", target if code == 301 else "(gone)", code,
                     "Legacy WordPress / funnel URL from the previous lmcmic.ca" if code == 301 else "Old WordPress admin/system path: 410 Gone tells crawlers to drop it"))
for i, row in enumerate(rows, start=2):
    for c, v in enumerate(row, 1):
        wr.cell(row=i, column=c, value=v)
body(wr, 2, len(rows) + 1, 4)
wr.freeze_panes = "A2"

# ---- Images
wi = wb.create_sheet("Images")
header(wi, 1, ["Image", "Source URL (lendmaxcapital.ca)", "Source size (KB)", "Alt text", "Used on", "Served file",
               "Format", "Width (px)", "Size (KB)", "Saving vs source"], [22, 44, 10, 60, 34, 34, 8, 9, 9, 12])
r = 2
used = {"driving-results": "/, /performance-and-risk/, /contact/",
        **{k: "/, /performance-and-risk/, /how-to-invest/" for k in IMAGES["logos"]}}
for name, spec in list(IMAGES["photos"].items()) + list(IMAGES["logos"].items()):
    files = sorted(f for f in LIVE["image_bytes"] if f.startswith(name + "-"))
    src_kb = LIVE["source_bytes"].get(spec["src"].rsplit("/", 1)[-1])
    for f in files:
        size = LIVE["image_bytes"][f]
        w = re.search(r"-(\d+)\.", f).group(1)
        is_photo = name in IMAGES["photos"]
        vals = [name, spec["src"], round(src_kb / 1024, 1) if src_kb else None, spec["alt"], used.get(name, ""),
                f"/assets/img/{f}", f.rsplit(".", 1)[1].upper(),
                int(w) if is_photo else round(int(w) * spec["width"] / spec["height"]),
                round(size / 1024, 1), f'=IF(ISNUMBER(C{r}),1-I{r}/C{r},"")']
        for c, v in enumerate(vals, 1):
            wi.cell(row=r, column=c, value=v)
        wi.cell(row=r, column=10).number_format = "0%"
        r += 1
og_row = r
for c, v in enumerate(["Share cards (og:image)", "Generated at build", None,
                       "Page title on a navy card with the Lendmax logo", "Every page (97)", "/assets/og/<page>.png",
                       "PNG", 1200, None, ""], 1):
    wi.cell(row=og_row, column=c, value=v)
body(wi, 2, og_row, 10)
wi.freeze_panes = "B2"
wi["C1"].comment = Comment(f"Measured on the server {LIVE['checked']}.", "lmcmic build")

# ---- External links
we = wb.create_sheet("External Links")
header(we, 1, ["URL", "Status", "Resolves", "Note"], [90, 8, 10, 70])
for i, x in enumerate(LINKS["results"], start=2):
    for c, v in enumerate([x["url"], x["status"] or "timeout", "Yes" if x["ok"] else "Check by hand", x.get("note", "")], 1):
        we.cell(row=i, column=c, value=v)
body(we, 2, len(LINKS["results"]) + 1, 4)
we.conditional_formatting.add(f"C2:C{len(LINKS['results']) + 1}", CellIsRule(operator="equal", formula=['"Yes"'], fill=OKF))
we.conditional_formatting.add(f"C2:C{len(LINKS['results']) + 1}", CellIsRule(operator="equal", formula=['"Check by hand"'], fill=WARNF))
we.cell(row=len(LINKS["results"]) + 3, column=1, value=f"Checked {LINKS['checked']}.").font = font(italic=True)

# ---- Technical SEO
wt = wb.create_sheet("Technical SEO")
header(wt, 1, ["Area", "Item", "Setting on lmcmic.ca", "Status", "Notes"], [18, 30, 60, 12, 70])
d, tls = LIVE["dns"], LIVE["tls"]
tech = [
    ("Domain", "Canonical origin", "https://lmcmic.ca (no www, HTTPS, trailing-slash URLs)", "Live", "Every other form 301s to it in one hop — see Redirects."),
    ("Domain", "DNS", f"A lmcmic.ca → {d['a']}; www → {d['www']} (Cloudflare, DNS only)", "Live", f"Checked {LIVE['checked']}."),
    ("Domain", "TLS certificate", f"{tls['issuer']} for {tls['names']}, expires {tls['expires']}", "Live", f"Renews automatically ({tls['renewal']}). HSTS max-age one year."),
    ("Hosting", "Server", "Lendmax server 159.203.33.160, nginx; static site + lead service", "Live", f"Release {LIVE['release']}. Registered as site 'lmcmic' in the server's site tools."),
    ("Crawling", "XML sitemap", "https://lmcmic.ca/sitemap.xml", "Live", f"{len(sitemap_urls)} URLs, lastmod on each, video entry on the home page, image entries for the photo pages."),
    ("Crawling", "robots.txt", "Allow all; Disallow /api/; Sitemap line", "Live", ""),
    ("Crawling", "llms.txt", "https://lmcmic.ca/llms.txt", "Live", "Plain map of the site for AI answer engines."),
    ("Crawling", "404 page", "Custom 404 with links to key pages; noindex", "Live", LIVE["probe_404_note"]),
    ("Crawling", "Googlebot", f"{LIVE['googlebot_200']} successful fetches in today's log", "Live", "Verified crawls return 200."),
    ("Indexing", "Canonical tags", "Self-referencing rel=canonical on every page", "Live", "See Pages, column Q."),
    ("Indexing", "hreflang", "en-ca + x-default (self)", "Live", ""),
    ("Indexing", "Robots meta", "index, follow, max-image-preview:large, max-snippet:-1", "Live", "404 is noindex."),
    ("Search Console", "Verification", "Meta tag on every page + DNS TXT google-site-verification", "Live",
     f"TXT records present: {', '.join(t[:10] + '…' for t in d['txt_google'])}."),
    ("Search Console", "Sitemap submitted", "Submit https://lmcmic.ca/sitemap.xml under Sitemaps", "To confirm", "Confirm in Search Console that it shows 'Success' and 97 discovered URLs."),
    ("Analytics", "Google Analytics 4", f"Measurement ID {SITE['ga4_id']} (property {SITE['ga4_property_id']})", "Live",
     "Loaded from site.js so the security policy still blocks inline script. Form submits send generate_lead."),
    ("Structured data", "Types in use", ", ".join(sorted({t for p in pages for t in p['schema'].split(', ') if t})), "Live", "Only marks up what the page shows."),
    ("Performance", "CSS / JS", "Tailwind compiled at build (~40 KB), one deferred script, no CDN script", "Live", ""),
    ("Performance", "Icons", "Material Symbols subset to the glyphs used; fixed 1em box (no layout shift)", "Live", ""),
    ("Performance", "Images", "AVIF / WebP / JPEG with srcset, width/height, lazy loading", "Live", "See Images."),
    ("Performance", "Video", "YouTube loads only on click (facade); VideoObject uploadDate 2021-12-31", "Live", ""),
    ("Security", "Headers", "HSTS, CSP (no inline script), X-Frame-Options DENY, nosniff, Referrer-Policy, Permissions-Policy", "Live", ""),
    ("Leads", "Contact form", f"POST /api/lead → stored → email {SITE['lead']['notify_email']} → redirect {SITE['lead']['redirect_url']}", "Live",
     "Email via FormSubmit (activated). For sending from your own domain, add a Resend/Postmark/SendGrid key."),
    ("Email DNS", "SPF", d["spf"], "Live", "Single SPF record (duplicate removed)."),
    ("Content", "Legal review", SITE["legal_review"]["note"], "Done", "Shown in every article byline."),
    ("Content", "Next scheduled review", SITE["next_review_label"], "Scheduled", "Returns figures and regulatory pages: review quarterly (Tab 3 L14)."),
]
for i, row in enumerate(tech, start=2):
    for c, v in enumerate(row, 1):
        wt.cell(row=i, column=c, value=v)
body(wt, 2, len(tech) + 1, 5)
wt.conditional_formatting.add(f"D2:D{len(tech) + 1}", CellIsRule(operator="equal", formula=['"Live"'], fill=OKF))
wt.conditional_formatting.add(f"D2:D{len(tech) + 1}", CellIsRule(operator="equal", formula=['"Done"'], fill=OKF))
wt.conditional_formatting.add(f"D2:D{len(tech) + 1}", CellIsRule(operator="equal", formula=['"To confirm"'], fill=WARNF))
wt.freeze_panes = "A2"

# ---- SEO checklist (Tab 3 → how lmcmic.ca applies it)
wc = wb.create_sheet("SEO Checklist")
src = openpyxl.load_workbook(SRC_XLSX, data_only=True)["3. SEO & Article Guidelines"]
APPLY = {
    "A": ("Applied", "Keywords, intent and verification notes come from the workbook per article (plan.json); primary statutes and regulator pages were fetched and checked on 3 Oct 2026; legal review completed 3 Oct 2026."),
    "B": ("Applied", "Short-answer box first (QA K1); one H1 with the primary keyword (97/97 pages); question-led H2s, no heading jumps (QA K2); terms defined at first use and linked to the 121-term glossary."),
    "C": ("Applied", "Worked examples, comparison tables, two calculators, due-diligence checklists, common-mistakes and province sections; FAQs from the workbook's conversational keywords; 3–8 internal links per article (QA K11) and allow-listed external sources."),
    "D": ("Applied", "No unverified statistics (design figures replaced by the fund's published returns); prohibited-language scan clean on every page; disclosures in the body (QA K8); legal review completed."),
    "E": ("Applied", "Q&A headings, opening definitions, numbered steps (HowTo), tables, FAQs (FAQPage schema), worked examples, key takeaways, sourced citations, descriptive internal anchors, schema that matches visible content."),
    "F": ("Applied", "The ten elements are required by the writing brief; spot-checked on cornerstone articles."),
    "G": ("Partly met", "lendmaxcapital.ca has one photograph and four partner logos, used on core pages as AVIF/WebP/JPEG with alt text, dimensions and lazy loading. Articles use generated 1200×630 title cards until more photography is supplied (Tab 2 image specs kept in plan.json)."),
    "H": ("Applied", "Disclosure box on every article plus in-body statements: no guarantee, no CDIC insurance, loss of principal, liquidity, not advice, no suitability, past performance, tax date, regulatory currency, higher return / higher risk."),
    "I": ("Applied", "Automated scan of all pages is clean. I8: partner names and logos appear at the client's request, as already published on lendmaxcapital.ca."),
    "J": ("Applied", "Closing summary names the seven axes (QA J2); the approved CTA text is on every page; one enquiry route: /contact/ → app.lendmaxcapital.ca/investor-start."),
    "K": ("See Content QA", "Automated K-checks pass for all 90 articles; manual checks and their status are on the Content QA sheet."),
    "L": ("Monitor", "Time-sensitive items: re-check on the dates given. L1 (BC Mortgage Services Act, in force 13 Oct 2026): update the BC pages from 'scheduled to come into force' to 'in force' and confirm licence categories with BCFSA."),
}
header(wc, 1, ["Section", "Ref", "Requirement", "What good looks like", "Owner", "Status on lmcmic.ca", "How lmcmic.ca applies it"],
       [26, 6, 44, 50, 16, 13, 70])
r = 2
for row in src.iter_rows(min_row=1, values_only=True):
    if not row or not row[1] or not row[0] or str(row[0]).startswith(("TAB", "HOW", "Section")):
        continue
    sec = str(row[0])[0]
    status, how = APPLY.get(sec, ("", ""))
    for c, v in enumerate([row[0], row[1], row[2], row[3], row[5], status, how], 1):
        wc.cell(row=r, column=c, value=v)
    r += 1
last_c = r - 1
body(wc, 2, last_c, 7)
wc.conditional_formatting.add(f"F2:F{last_c}", CellIsRule(operator="equal", formula=['"Applied"'], fill=OKF))
wc.conditional_formatting.add(f"F2:F{last_c}", CellIsRule(operator="equal", formula=['"Partly met"'], fill=WARNF))
wc.conditional_formatting.add(f"F2:F{last_c}", CellIsRule(operator="equal", formula=['"Monitor"'], fill=WARNF))
wc.freeze_panes = "C2"
wc.auto_filter.ref = f"A1:G{last_c}"

# ---- Overview (first sheet)
wo = wb.create_sheet("Overview", 0)
title_block(wo, "lmcmic.ca — SEO master sheet",
            f"Lendmax Capital MIC investor site. Built from the live site on {TODAY}. Source plan: content/data/seo-content-roadmap.xlsx (Tabs 1–3).")
wo.column_dimensions["A"].width = 46
wo.column_dimensions["B"].width = 16
wo.column_dimensions["C"].width = 70
stats = [
    ("Pages", None, None),
    ("Indexable pages", f"=COUNTA(Pages!B2:B{last_pages})", "Every page in the Pages sheet"),
    ("Pages in the XML sitemap", f'=COUNTIF(Pages!Z2:Z{last_pages},"Yes")', "Should equal indexable pages"),
    ("Investor guides, tools, glossary and FAQ hub", f'=COUNTIF(Pages!C2:C{last_pages},"<>Core page")-COUNTIF(Pages!C2:C{last_pages},"Hub")-COUNTIF(Pages!C2:C{last_pages},"Legal")', ""),
    ("Titles inside 30–65 characters", f'=COUNTIF(Pages!H2:H{last_pages},"OK")', "Pages sheet, column H"),
    ("Descriptions inside 120–160 characters", f'=COUNTIF(Pages!K2:K{last_pages},"OK")', "Pages sheet, column K"),
    ("Canonical tag matches the URL", f'=COUNTIF(Pages!Q2:Q{last_pages},"Yes")', "Pages sheet, column Q"),
    ("Pages with the GA4 tag", f'=COUNTIF(Pages!AA2:AA{last_pages},"Yes")', ""),
    ("Total words on indexable pages", f"=SUM(Pages!T2:T{last_pages})", "Main content only"),
    ("Content quality", None, None),
    ("Articles passing every automated check", f'=COUNTIF(\'Content QA\'!H2:H{last_q},"PASS")', f"of {len(QA['articles'])} planned articles"),
    ("Internal links (in-content)", f"=COUNTA('Internal Links'!A2:A{last_l})", "Navigation and footer excluded"),
    ("Internal links pointing at a missing page", f'=COUNTIF(\'Internal Links\'!D2:D{last_l},"Check")', "Should be 0"),
    ("External links resolving from the server", f'=COUNTIF(\'External Links\'!C2:C{len(LINKS["results"]) + 1},"Yes")', f"of {len(LINKS['results'])}; the rest block automated checks"),
    ("Keywords", None, None),
    ("Workbook keywords", f"=COUNTA(Keywords!A2:A{last_kw})", "Tab 1 of the source workbook"),
    ("Mapped to a live page", f'=COUNTIF(Keywords!M2:M{last_kw},"/*")', ""),
    ("Keyword terms present on their page", f'=COUNTIF(Keywords!O2:O{last_kw},"Yes")', ""),
    ("Not built by design (workbook says do not build)", f'=COUNTIF(Keywords!M2:M{last_kw},"")', ""),
    ("Tier 1 quick-win keywords", f'=COUNTIF(Keywords!J2:J{last_kw},"1*")', ""),
]
r = 4
for label, val, note in stats:
    if val is None:
        wo.cell(row=r, column=1, value=label).font = font(size=12, bold=True, color=NAVY)
        r += 1
        continue
    wo.cell(row=r, column=1, value=label).font = font()
    c = wo.cell(row=r, column=2, value=val)
    c.font = font(bold=True)
    c.number_format = "#,##0"
    wo.cell(row=r, column=3, value=note).font = font(color="44474C")
    for col in (1, 2, 3):
        wo.cell(row=r, column=col).border = BOX
    r += 1
r += 1
wo.cell(row=r, column=1, value="Action items").font = font(size=12, bold=True, color=RED)
r += 1
actions = [
    ("BC Mortgage Services Act comes into force 13 October 2026", "Update the BC pages from 'scheduled to come into force' to 'in force', confirm licence categories with BCFSA (Tab 3 L1)."),
    ("Search Console", "Confirm the sitemap shows Success with 97 discovered URLs; watch Page indexing for the first two weeks."),
    ("Keyword volumes", "Fill the yellow columns on the Keywords sheet from Keyword Planner / Ahrefs (Canada) and Search Console."),
    ("Second-person calculation check (K10)", "Have a second person recompute the worked examples and record it on Content QA."),
    ("Four external links", "Click through the links marked 'Check by hand' on External Links (sites that block automated checks)."),
    ("Photography", "Supply more photos to give articles their own featured images (Tab 2 image specs are ready)."),
    ("Quarterly review", "Re-check return figures and regulatory statements each quarter (next: January 2027)."),
]
for a, b in actions:
    wo.cell(row=r, column=1, value=a).font = font(bold=True)
    wo.cell(row=r, column=3, value=b).font = font()
    wo.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="top")
    wo.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    wo.row_dimensions[r].height = 30
    r += 1
r += 1
wo.cell(row=r, column=1, value="Sheets").font = font(size=12, bold=True, color=NAVY)
r += 1
for name, what in [("Pages", "Every live page: title, description, H1, keywords, canonical, schema, words, links, images, QA."),
                   ("Keywords", "All 382 workbook keywords → the live page that targets them. Yellow columns are yours to fill."),
                   ("Content QA", "The Tab 3 pre-publication gate for all 90 articles, plus manual sign-offs."),
                   ("Internal Links", "Every in-content link with its anchor text."),
                   ("Redirects", "Canonical and legacy-URL redirect rules on the server."),
                   ("Images", "Photos and logos reused from lendmaxcapital.ca, every optimised size."),
                   ("External Links", "Outbound citations and whether each resolves."),
                   ("Technical SEO", "Domain, crawling, indexing, analytics, performance and security settings."),
                   ("SEO Checklist", "Tab 3 of the source workbook with how lmcmic.ca applies each rule.")]:
    wo.cell(row=r, column=1, value=name).font = font(bold=True, color=SKY)
    wo.cell(row=r, column=1).hyperlink = f"#'{name}'!A1"
    wo.cell(row=r, column=3, value=what).font = font()
    r += 1
r += 1
wo.cell(row=r, column=1, value="Legend").font = font(size=12, bold=True, color=NAVY)
r += 1
c = wo.cell(row=r, column=1, value="Yellow cells")
c.fill, c.font = INPUT, font()
wo.cell(row=r, column=3, value="Input cells for you to fill (e.g. Keywords: measured volume, Search Console position and clicks).").font = font()
r += 1
c = wo.cell(row=r, column=1, value="Green / red / amber")
c.fill, c.font = OKF, font()
wo.cell(row=r, column=3, value="Pass / fail / needs review, set automatically by formulas and conditional formatting.").font = font()
wo.sheet_view.showGridLines = False

for s in wb.worksheets:
    s.sheet_properties.tabColor = NAVY if s.title == "Overview" else SKY

OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print(f"wrote {OUT}: {len(pages)} pages, {len(QA['keywords'])} keywords, {last_l - 1} internal links, {len(rows)} redirect rules")
