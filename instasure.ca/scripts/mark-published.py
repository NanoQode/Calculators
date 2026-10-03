"""Mark blog topics as published in docs/blog-content-plan.csv and the SEO workbook (Blog Content Plan tab).

Usage: python3 scripts/mark-published.py <slug> [<slug> …] [--writer "Michael Le Chi"] [--date YYYY-MM-DD]
Needs openpyxl (pip install openpyxl). Formulas and formatting in the workbook are kept.
"""
import csv, datetime, os, sys
from openpyxl import load_workbook

ROOT = os.path.join(os.path.dirname(__file__), '..')
CSV = os.path.join(ROOT, 'docs', 'blog-content-plan.csv')
XLSX = os.path.join(ROOT, 'docs', 'instasure-seo-keyword-content-plan.xlsx')
args = sys.argv[1:]
def opt(name, default):
    if name in args:
        i = args.index(name); v = args[i + 1]; del args[i:i + 2]; return v
    return default
writer = opt('--writer', 'Michael Le Chi')
date = opt('--date', datetime.date.today().isoformat())
slugs = set(args)
if not slugs: sys.exit('give at least one slug')
url = lambda s: f'https://instasure.ca/guides/{s}/'

rows = list(csv.DictReader(open(CSV, encoding='utf-8')))
fields = list(rows[0].keys())
hit = set()
for r in rows:
    for s in slugs:
        if r['URL'] == url(s): r['Status'], r['Writer'], r['Published date'] = 'Published', writer, date; hit.add(s)
with open(CSV, 'w', encoding='utf-8', newline='') as f:
    w = csv.DictWriter(f, fieldnames=fields); w.writeheader(); w.writerows(rows)

wb = load_workbook(XLSX)
ws = wb['Blog Content Plan']
head = [c.value for c in ws[1]]
col = {h: i + 1 for i, h in enumerate(head)}
for r in range(2, ws.max_row + 1):
    u = ws.cell(row=r, column=col['URL']).value
    for s in slugs:
        if u == url(s):
            ws.cell(row=r, column=col['Status'], value='Published')
            ws.cell(row=r, column=col['Writer'], value=writer)
            ws.cell(row=r, column=col['Published date'], value=date)
wb.save(XLSX)
missing = slugs - hit
print(f'marked {len(hit)} published ({date}, {writer})' + (f'; not in the plan: {", ".join(sorted(missing))}' if missing else ''))
